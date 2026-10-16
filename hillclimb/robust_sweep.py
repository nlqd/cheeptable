#!/usr/bin/env python3
"""Robustness sweep: degrade the two real pages (3-4 of the contract) and see where the grid pipeline breaks.
  robust_sweep.py OUT.json        run with ../venv_onnx_export/bin/python and GRID_V2=1
Each condition: render, degrade the page image, run the ruled-grid reader, score rows found / numbers / description error
against the corrected golden. The python prototype is used because it is fast; the JS app uses the same grid and recogniser."""
import json
import os
import sys
import tempfile

import cv2
import numpy as np
import openpyxl

sys.path.insert(0, "..")
sys.path.insert(0, ".")
import grid_extract  # noqa: E402
import score  # noqa: E402

PDF = "../7. Biểu giá HĐ.pdf"
rng = np.random.default_rng(0)


def rotate(deg):
    def f(img):
        h, w = img.shape
        m = cv2.getRotationMatrix2D((w / 2, h / 2), deg, 1.0)
        return cv2.warpAffine(img, m, (w, h), flags=cv2.INTER_LINEAR, borderValue=255)
    return f


def blur(sigma):
    return lambda img: cv2.GaussianBlur(img, (0, 0), sigma)


def lowres(factor):
    def f(img):
        h, w = img.shape
        small = cv2.resize(img, (int(w * factor), int(h * factor)), interpolation=cv2.INTER_AREA)
        return cv2.resize(small, (w, h), interpolation=cv2.INTER_LINEAR)
    return f


def jpeg(q):
    def f(img):
        ok, buf = cv2.imencode(".jpg", img, [cv2.IMWRITE_JPEG_QUALITY, q])
        return cv2.imdecode(buf, cv2.IMREAD_GRAYSCALE)
    return f


def contrast(alpha):
    return lambda img: np.clip(255 - (255 - img.astype(np.float32)) * alpha, 0, 255).astype(np.uint8)


def shadow(img):
    h, w = img.shape
    ramp = np.linspace(1.0, 0.55, w, dtype=np.float32)[None, :]
    return np.clip(img.astype(np.float32) * ramp, 0, 255).astype(np.uint8)


def noise(sigma):
    return lambda img: np.clip(img.astype(np.float32) + rng.normal(0, sigma, img.shape), 0, 255).astype(np.uint8)


CONDITIONS = [("baseline", None), ("rotate 0.5deg", rotate(0.5)), ("rotate 1deg", rotate(1)), ("rotate 2deg", rotate(2)),
              ("rotate 3deg", rotate(3)), ("blur s1", blur(1)), ("blur s2", blur(2)), ("resolution x0.75", lowres(0.75)),
              ("resolution x0.5", lowres(0.5)), ("resolution x0.35", lowres(0.35)), ("jpeg q50", jpeg(50)), ("jpeg q15", jpeg(15)),
              ("contrast x0.5", contrast(0.5)), ("contrast x0.3", contrast(0.3)), ("shadow gradient", shadow), ("noise s15", noise(15)),
              ("noise s30", noise(30))]

gt = openpyxl.Workbook()
ws = gt.active
for row in openpyxl.load_workbook("golden_v2.xlsx", data_only=True).active.iter_rows(values_only=True):
    try:
        if 94 <= int(str(row[0]).strip()) <= 188:
            ws.append(row[:6])
    except ValueError:
        pass
gt_path = os.path.join(tempfile.mkdtemp(), "gold.xlsx")
gt.save(gt_path)

results = []
for name, fn in CONDITIONS:
    rows = grid_extract.extract(PDF, {3, 4}, transform=fn)
    path = gt_path + ".rows.json"
    json.dump(rows, open(path, "w"))
    try:
        r = score.report(path, gt_path)["all"]
    except Exception as e:                       # no rows at all
        r = {"found": 0, "rows": 95, "all3": 0.0, "cer": 1.0, "desc_exact": 0.0}
    cols = len(rows[0]["cells"]) if rows else 0
    results.append({"condition": name, "rows_out": len(rows), "cols": cols, "found": r["found"], "all3": round(r["all3"], 3),
                    "exact": round(r["desc_exact"], 3), "cer": round(r["cer"], 3)})
    print("%-18s rows_out %3d cols %d  found %2d/95  all3 %.3f  exact %.3f  cer %.3f" % (
        name, len(rows), cols, r["found"], r["all3"], r["desc_exact"], r["cer"]), flush=True)
json.dump(results, open(sys.argv[1], "w"), indent=1)
